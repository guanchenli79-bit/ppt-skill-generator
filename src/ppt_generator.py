"""PowerPoint rendering with python-pptx."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, Iterable, List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .outline_generator import Slide, normalize_style


THEMES: Dict[str, Dict[str, str]] = {
    "商务风": {
        "background": "F7F8FA",
        "surface": "FFFFFF",
        "text": "1F2937",
        "muted": "667085",
        "accent": "1D4ED8",
        "accent_2": "111827",
        "line": "D0D5DD",
    },
    "科技风": {
        "background": "07111F",
        "surface": "102033",
        "text": "F8FAFC",
        "muted": "B6C2D0",
        "accent": "38BDF8",
        "accent_2": "5EEAD4",
        "line": "24435F",
    },
    "教育风": {
        "background": "F5F8FF",
        "surface": "FFFFFF",
        "text": "1E293B",
        "muted": "64748B",
        "accent": "2563EB",
        "accent_2": "F59E0B",
        "line": "CBD5E1",
    },
    "营销风": {
        "background": "FFF7ED",
        "surface": "FFFFFF",
        "text": "241C15",
        "muted": "7C6F64",
        "accent": "EA580C",
        "accent_2": "DB2777",
        "line": "FED7AA",
    },
    "极简风": {
        "background": "FAFAFA",
        "surface": "FFFFFF",
        "text": "18181B",
        "muted": "71717A",
        "accent": "27272A",
        "accent_2": "A1A1AA",
        "line": "E4E4E7",
    },
}

SLIDE_W = 13.333
SLIDE_H = 7.5


def generate_ppt(
    slides: List[Slide],
    output_path: str | Path,
    topic: str,
    style: str = "商务风",
    audience: str = "通用受众",
    purpose: str = "演示汇报",
) -> Path:
    """Render a list of slides to a real PPTX file."""

    normalized_style = normalize_style(style)
    theme = THEMES.get(normalized_style, THEMES["商务风"])
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for index, slide_spec in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _paint_background(slide, theme)
        _add_page_marker(slide, theme, index, len(slides))

        if slide_spec.slide_type == "cover":
            _render_cover(slide, slide_spec, theme, topic, audience, purpose)
        elif slide_spec.slide_type == "agenda":
            _render_agenda(slide, slide_spec, theme)
        elif slide_spec.slide_type == "summary":
            _render_summary(slide, slide_spec, theme)
        elif slide_spec.slide_type == "closing":
            _render_closing(slide, slide_spec, theme)
        else:
            _render_content(slide, slide_spec, theme, index)

    prs.save(output)
    return output


def default_output_path(topic: str) -> Path:
    """Create a stable default output path from the topic."""

    slug = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff]+", "-", topic.strip()).strip("-")
    slug = slug or "presentation"
    return Path("output") / f"{slug}.pptx"


def _paint_background(slide, theme: Dict[str, str]) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(theme["background"])

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(SLIDE_H))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(theme["accent"])
    accent.line.fill.background()


def _render_cover(slide, spec: Slide, theme: Dict[str, str], topic: str, audience: str, purpose: str) -> None:
    _add_text(slide, purpose, 0.75, 0.8, 6.5, 0.35, 12, theme["accent"], bold=True)
    _add_text(slide, topic or spec.title, 0.75, 1.45, 7.6, 1.2, 34, theme["text"], bold=True)
    _add_text(slide, f"面向：{audience}", 0.78, 2.85, 5.6, 0.4, 15, theme["muted"])

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(1.1), Inches(3.8), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb(theme["surface"])
    card.line.color.rgb = _rgb(theme["line"])

    _add_text(slide, "PPT Skill Generator", 8.72, 1.55, 3.1, 0.35, 15, theme["accent"], bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "\n".join(spec.bullets), 8.85, 2.35, 2.8, 1.6, 13, theme["text"], align=PP_ALIGN.CENTER)
    _add_badge(slide, theme, 8.95, 4.55, "Editable PPTX")


def _render_agenda(slide, spec: Slide, theme: Dict[str, str]) -> None:
    _add_title(slide, spec.title, spec.subtitle, theme)
    y = 1.85
    for index, item in enumerate(spec.bullets, start=1):
        _add_number(slide, theme, 0.92, y, index)
        _add_text(slide, item, 1.55, y + 0.02, 8.8, 0.35, 18, theme["text"], bold=index == 1)
        y += 0.78


def _render_content(slide, spec: Slide, theme: Dict[str, str], index: int) -> None:
    _add_title(slide, spec.title, spec.subtitle, theme)

    x_positions = [0.85, 6.9] if len(spec.bullets) > 2 else [1.1]
    for bullet_index, bullet in enumerate(spec.bullets):
        col = bullet_index % len(x_positions)
        row = bullet_index // len(x_positions)
        x = x_positions[col]
        y = 2.0 + row * 1.25
        w = 5.45 if len(x_positions) > 1 else 10.8
        _add_bullet_card(slide, theme, x, y, w, 0.9, bullet_index + 1, bullet)

    if index % 3 == 0:
        _add_side_metric(slide, theme, "Key", "Focus")


def _render_summary(slide, spec: Slide, theme: Dict[str, str]) -> None:
    _add_title(slide, spec.title, spec.subtitle, theme)
    for index, bullet in enumerate(spec.bullets, start=1):
        x = 0.9 + (index - 1) * 4.05
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.25), Inches(3.35), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(theme["surface"])
        card.line.color.rgb = _rgb(theme["line"])
        _add_text(slide, f"0{index}", x + 0.28, 2.58, 0.8, 0.25, 14, theme["accent"], bold=True)
        _add_text(slide, bullet, x + 0.28, 3.15, 2.75, 0.85, 15, theme["text"], bold=True)


def _render_closing(slide, spec: Slide, theme: Dict[str, str]) -> None:
    _add_text(slide, spec.title, 2.1, 1.75, 9.2, 0.8, 42, theme["text"], bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, spec.subtitle, 2.3, 2.85, 8.8, 0.35, 16, theme["muted"], align=PP_ALIGN.CENTER)
    _add_badge(slide, theme, 4.95, 4.2, "Generated by PPT Skill Generator", width=3.45)


def _add_title(slide, title: str, subtitle: str, theme: Dict[str, str]) -> None:
    _add_text(slide, title, 0.85, 0.72, 9.6, 0.55, 27, theme["text"], bold=True)
    _add_text(slide, subtitle, 0.88, 1.35, 9.6, 0.35, 12.5, theme["muted"])


def _add_bullet_card(slide, theme: Dict[str, str], x: float, y: float, w: float, h: float, index: int, text: str) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = _rgb(theme["surface"])
    card.line.color.rgb = _rgb(theme["line"])

    _add_number(slide, theme, x + 0.22, y + 0.22, index, size=0.36)
    _add_text(slide, text, x + 0.78, y + 0.17, w - 1.05, h - 0.25, 12.8, theme["text"])


def _add_number(slide, theme: Dict[str, str], x: float, y: float, index: int, size: float = 0.42) -> None:
    marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    marker.fill.solid()
    marker.fill.fore_color.rgb = _rgb(theme["accent"])
    marker.line.fill.background()
    _add_text(slide, str(index), x, y + 0.08, size, 0.12, 8, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)


def _add_side_metric(slide, theme: Dict[str, str], label: str, value: str) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.75), Inches(5.15), Inches(1.65), Inches(0.72))
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(theme["accent"])
    box.line.fill.background()
    _add_text(slide, label, 10.92, 5.28, 0.5, 0.15, 8, "FFFFFF", bold=True)
    _add_text(slide, value, 11.45, 5.25, 0.7, 0.2, 12, "FFFFFF", bold=True, align=PP_ALIGN.RIGHT)


def _add_badge(slide, theme: Dict[str, str], x: float, y: float, text: str, width: float = 2.1) -> None:
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(0.48))
    badge.fill.solid()
    badge.fill.fore_color.rgb = _rgb(theme["accent"])
    badge.line.fill.background()
    _add_text(slide, text, x + 0.14, y + 0.15, width - 0.28, 0.13, 9, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)


def _add_page_marker(slide, theme: Dict[str, str], page: int, total: int) -> None:
    _add_text(slide, f"{page}/{total}", 11.55, 6.95, 0.75, 0.18, 8, theme["muted"], align=PP_ALIGN.RIGHT)


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float,
    color: str,
    bold: bool = False,
    align: PP_ALIGN | None = None,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True

    paragraph = frame.paragraphs[0]
    paragraph.alignment = align or PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
